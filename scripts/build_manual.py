"""Build the moreader operator manual as a Word (.docx) file.

Run:  python3 scripts/build_manual.py /tmp/manual docs/moreader-operator-manual.docx
"""

import sys
from datetime import date
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt, RGBColor

SHOTS = Path(sys.argv[1] if len(sys.argv) > 1 else "/tmp/manual")
OUT = Path(sys.argv[2] if len(sys.argv) > 2 else "docs/moreader-operator-manual.docx")
OUT.parent.mkdir(parents=True, exist_ok=True)

NAVY = RGBColor(0x1F, 0x2D, 0x3D)
BLUE = RGBColor(0x2D, 0x8C, 0xF0)
GREEN = RGBColor(0x15, 0x80, 0x3D)
RED = RGBColor(0xB9, 0x1C, 0x1C)
AMBER = RGBColor(0xB7, 0x79, 0x1F)

doc = Document()

# Base font.
normal = doc.styles["Normal"]
normal.font.name = "Calibri"
normal.font.size = Pt(12)
doc.styles["Heading 1"].font.color.rgb = NAVY
doc.styles["Heading 2"].font.color.rgb = BLUE


def para(text="", size=12, bold=False, italic=False, color=None, align=None, space_after=6):
    p = doc.add_paragraph()
    if align is not None:
        p.alignment = align
    p.paragraph_format.space_after = Pt(space_after)
    if text:
        r = p.add_run(text)
        r.bold = bold
        r.italic = italic
        r.font.size = Pt(size)
        if color is not None:
            r.font.color.rgb = color
    return p


def bullet(text):
    doc.add_paragraph(text, style="List Bullet")


def step(text):
    doc.add_paragraph(text, style="List Number")


def picture(name, caption, width=6.5):
    path = SHOTS / name
    if not path.exists():
        para(f"[missing image: {name}]", italic=True, color=RED)
        return
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run().add_picture(str(path), width=Inches(width))
    cap = para(caption, size=10, italic=True, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=12)
    cap.runs[0].font.color.rgb = RGBColor(0x60, 0x60, 0x60)


def heading(text, level=1):
    doc.add_heading(text, level=level)


# ---------------------------------------------------------------- title page
t = para("moreader", size=40, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=0)
t.runs[0].font.color.rgb = NAVY
para("Operator Manual", size=22, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, color=BLUE)
para("Manufacturing Order Verification System", size=14, align=WD_ALIGN_PARAGRAPH.CENTER)
para("")
para("This guide shows you how to use the moreader screen on the line.",
     size=13, align=WD_ALIGN_PARAGRAPH.CENTER)
para("")
para(f"Version 1.0   ·   {date.today().strftime('%B %Y')}",
     size=11, align=WD_ALIGN_PARAGRAPH.CENTER, color=RGBColor(0x60, 0x60, 0x60))
doc.add_page_break()

# ---------------------------------------------------------------- contents
heading("What Is in This Guide", 1)
for item in [
    "1. What This System Does",
    "2. The Main Screen",
    "3. Colors and Symbols",
    "4. How to Scan a Manufacturing Order",
    "5. When a Scan Does Not Pass",
    "6. The Bypass Button",
    "7. The Manual Lockout Button",
    "8. Shift Change",
    "9. The Heartbeat Light",
    "10. Simple Problems and Fixes",
    "11. Safety",
    "12. Who to Call",
]:
    para(item, size=12, space_after=2)
doc.add_page_break()

# ---------------------------------------------------------------- 1
heading("1. What This System Does", 1)
para("This system makes sure the line runs the right product. It checks the "
     "manufacturing order (MO) before the line can run.")
para("An MO is the paper order. It tells the line what to build. Each MO has a "
     "barcode.")
para("Here is the simple idea:")
bullet("You scan the MO barcode.")
bullet("The system reads the last 4 numbers on the barcode.")
bullet("It checks those numbers against each machine (the encapsulators).")
bullet("If they all match, the line is allowed to run.")
bullet("If they do not match, the line stays stopped and shows you why.")
para("The line cannot run until the MO passes. This stops the wrong product from "
     "being built.")

# ---------------------------------------------------------------- 2
heading("2. The Main Screen", 1)
para("This is the screen you use most. It shows three machine boxes on top and "
     "the master box (COS) below them.")
picture("01_main_locked.png", "Figure 1. The main screen, waiting for a scan.")
para("Here is what each part means:", bold=True)
bullet("Machine boxes (Encapsulator 1, 2, 3): each box shows one machine.")
bullet("RECIPE (DINT): the big number is the product the machine is set to run.")
bullet("LAST SCAN: the number from the last MO you scanned.")
bullet("COS · MASTER: the master box. It shows if the line is allowed to run.")
bullet("VERIFY MO (blue button): press this to start a scan.")
bullet("BYPASS (purple button): a supervisor-only way to skip the check.")
bullet("MANUAL LOCKOUT (red button): press this to stop the line and require a new scan.")
bullet("Top bar: shows how many machines are online, the time, and Settings.")
bullet("Bottom box: a running list of what just happened.")

# ---------------------------------------------------------------- 3
heading("3. Colors and Symbols", 1)
para("The screen uses colors and symbols so you can tell the state fast.")
picture("03_tile_states.png", "Figure 2. Left: good. Middle: bad. Right: waiting.", 6.5)
para("Green check mark = GOOD.", bold=True, color=GREEN)
para("A green box with a check mark (✓) means this machine matches the MO.")
para("Red X = BAD.", bold=True, color=RED)
para("A red box with an X (✗) means this machine does not match the MO. The line "
     "will not run.")
para("Amber (yellow) = WAITING.", bold=True, color=AMBER)
para("An amber box means the machine is waiting for a scan.")
para("The master box (COS) at the bottom:", bold=True)
bullet("Green “MO VERIFIED” means the line is allowed to run.")
bullet("Amber “LOCKED — SCAN MO” means you must scan first.")
bullet("Purple “MO BYPASSED” means a supervisor turned on bypass.")

# ---------------------------------------------------------------- 4
heading("4. How to Scan a Manufacturing Order", 1)
para("Do this at the start of your shift, or any time the screen says "
     "“LOCKED — SCAN MO.”")
para("Your line uses three scans. You do them in order. A pop-up opens and "
     "guides you. It shows which step you are on, like “Step 1 of 3.”")
para("To start, press the blue VERIFY MO button.", bold=True)
para("You never type a number by hand. You only scan.")

heading("Step 1: Scan the Stuffed Element MO", 2)
para("Scan the Stuffed Element MO barcode. In this example it is 2220-1321. "
     "The system uses the last 4 numbers (1321) to check each machine.")
picture("08_scan_step1.png", "Figure 3. Step 1 of 3 — scan the Stuffed Element MO (2220-1321).")

heading("Step 2: Scan the Assembled Battery MO", 2)
para("Next, scan the Assembled Battery MO barcode. In this example it is "
     "2220-1964. This is a different order than Step 1.")
picture("09_scan_step2.png", "Figure 4. Step 2 of 3 — scan the Assembled Battery MO (2220-1964).")

heading("Step 3: Scan the Battery Label", 2)
para("Last, scan the battery label barcode. In this example it is "
     "1964W1261820308. The first 4 numbers (1964) must match the Assembled "
     "Battery MO from Step 2.")
picture("10_scan_step3.png", "Figure 5. Step 3 of 3 — scan the battery label (1964W1261820308).")

heading("When You Are Done", 2)
para("After the last scan, press OK. Some scanners do this for you.")
para("If everything matches, the machine boxes turn green. The master box says "
     "“MO VERIFIED.” The line can now run.")
picture("02_main_verified.png", "Figure 6. All three scans passed. The line is ready.")
para("Remember:", bold=True)
bullet("Scan in order: Stuffed Element, then Assembled Battery, then battery label.")
bullet("You only scan. You do not type any number.")
bullet("If any scan does not match, a red screen tells you what is wrong. See Section 5.")

# ---------------------------------------------------------------- 5
heading("5. When a Scan Does Not Pass", 1)
para("If a scan does not pass, a red screen pops up. It tells you what is wrong.")
picture("05_error_screen.png", "Figure 7. The error screen shows the reason.", 5.5)
para("What to do:", bold=True)
step("Read the reason on the red screen.")
step("Press OK.")
step("Fix the problem. For example, get the right MO, or check the machine setup.")
step("Scan again.")
para("The line stays stopped until a scan passes. This is normal. It keeps the "
     "wrong product from being built.")

# ---------------------------------------------------------------- 6
heading("6. The Bypass Button", 1)
para("Bypass lets the line run without a matching scan. It is for supervisors "
     "only. Use it only when you are told to.")
step("Press the purple BYPASS button.")
step("A box asks for a password.")
step("Type the password and press OK.")
picture("06_bypass_password.png", "Figure 8. Bypass asks for a password.", 4.6)
para("When bypass is on, the master box turns purple and says “MO BYPASSED.”")
picture("07_bypass_active.png", "Figure 9. Bypass is on. The master box is purple.")
para("Important:", bold=True, color=RED)
bullet("Bypass turns off by itself at a shift change.")
bullet("Bypass also turns off when someone presses MANUAL LOCKOUT.")
bullet("Every bypass is saved in the record with the date and time.")

# ---------------------------------------------------------------- 7
heading("7. The Manual Lockout Button", 1)
para("Press the red MANUAL LOCKOUT button to stop the line on purpose.")
para("When you press it:")
bullet("The line is no longer allowed to run.")
bullet("The machine boxes turn amber and say “LOCKED — SCAN MO.”")
bullet("You must scan the MO again to start.")
para("Use this when you need a fresh scan, or when you are told to lock the line.")

# ---------------------------------------------------------------- 8
heading("8. Shift Change", 1)
para("At each shift change, the system locks the line by itself. The screen will "
     "say “LOCKED — SCAN MO.”")
para("This means the new shift must scan the MO before the line can run. Just "
     "follow the steps in Section 4.")

# ---------------------------------------------------------------- 9
heading("9. The Heartbeat Light", 1)
para("The master box has a small heart symbol. It flips between ON and OFF over "
     "and over. This tells the PLC that the program is alive and working.")
bullet("If the heart keeps changing, the program is healthy.")
bullet("If the heart stops changing, tell maintenance.")

# ---------------------------------------------------------------- 10
heading("10. Simple Problems and Fixes", 1)
para("A machine box is gray and says “OFFLINE.”", bold=True)
bullet("This machine is not talking to the system.")
bullet("Check that the machine is powered on and on the network. Tell maintenance if it stays gray.")
para("The top bar does not say “PLCs online: 4/4.”", bold=True)
bullet("One or more PLCs are offline. The line will not verify until they are back.")
para("The scanner does not read.", bold=True)
bullet("Check the scanner cable. Try scanning again, straight and close.")
bullet("Make sure you pressed VERIFY MO first so the pop-up is open.")
para("The scan will not pass, but the MO looks right.", bold=True)
bullet("Read the red error screen. It tells you exactly what did not match.")
bullet("Check the machine recipe number against the MO.")

# ---------------------------------------------------------------- 11
heading("11. Safety", 1)
para("This system helps stop the wrong product from being built. It is not a "
     "safety guard.")
bullet("Always follow all machine safety rules.")
bullet("Never reach into a running machine.")
bullet("E-stops, guards, and safety devices work on their own. Do not rely on this screen for safety.")

# ---------------------------------------------------------------- 12
heading("12. Who to Call", 1)
para("Write your site contacts here:")
para("Supervisor: ______________________________")
para("Maintenance: _____________________________")
para("Controls / IT: ___________________________")

doc.save(str(OUT))
print("wrote", OUT)
